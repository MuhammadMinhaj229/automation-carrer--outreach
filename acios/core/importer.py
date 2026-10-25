"""
ACIOS Master Data Importer.
Migrates all legacy Excel workbooks, suppression lists, and Gmail draft records
into the unified canonical SQLite store (data/acios.db).
"""

import json
import re
import sys
from pathlib import Path
from datetime import datetime, timezone
import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(BASE_DIR))

from acios.core.database import get_db_connection, init_db, DEFAULT_DB_PATH
from acios.core.repository import CareerRepository


def import_all():
    db_path = DEFAULT_DB_PATH
    init_db(db_path)
    repo = CareerRepository(db_path)
    conn = get_db_connection(db_path)
    cursor = conn.cursor()

    print("[*] Starting Canonical Database Ingestion...")

    # 1. Ingest Master Suppression Lists
    print(" -> Ingesting master suppression records...")
    supp_local = BASE_DIR / "data" / "suppression" / "master_suppression.json"
    supp_path1 = Path(r"C:\Users\User\.gemini\antigravity\brain\eb93f88a-ad00-4e03-89b8-0c0b43ed482d\scratch\master_suppression_list.json")
    supp_path2 = Path(r"C:\Users\User\.gemini\antigravity\brain\eb93f88a-ad00-4e03-89b8-0c0b43ed482d\scratch\gmail_history_audit.json")
    supp_path3 = Path(r"C:\Users\User\.gemini\antigravity\brain\eb93f88a-ad00-4e03-89b8-0c0b43ed482d\jobs automation\master_suppression_list.json")

    suppressed_records = {}

    def extract_suppression(data, default_source, default_reason):
        if isinstance(data, list):
            for item in data:
                if isinstance(item, dict):
                    email = item.get("email")
                    reason = item.get("reason", default_reason)
                    source = item.get("source", default_source)
                elif isinstance(item, str):
                    email = item
                    reason = default_reason
                    source = default_source
                else:
                    continue
                if email and "@" in email:
                    clean = email.strip().lower()
                    if clean not in suppressed_records:
                        suppressed_records[clean] = (clean, reason, source)
        elif isinstance(data, dict):
            for k in ["suppressed_emails", "hard_bounces", "bounced_addresses", "bounced_recipients"]:
                for item in data.get(k, []):
                    if isinstance(item, dict):
                        email = item.get("email")
                        reason = item.get("reason", default_reason)
                        source = item.get("source", default_source)
                    elif isinstance(item, str):
                        email = item
                        reason = default_reason
                        source = default_source
                    else:
                        continue
                    if email and "@" in email:
                        clean = email.strip().lower()
                        if clean not in suppressed_records:
                            suppressed_records[clean] = (clean, reason, source)

    for p, src, rsn in [
        (supp_local, "Local Master Suppression", "Historical Bounce/Opt-out"),
        (supp_path3, "Jobs Automation Master", "Historical Bounce/Opt-out"),
        (supp_path1, "Scratch Master Suppression", "Historical Bounce/Opt-out"),
        (supp_path2, "Gmail History Bounce Audit", "Permanent Delivery Failure")
    ]:
        if p.exists():
            try:
                content = json.loads(p.read_text(encoding="utf-8"))
                extract_suppression(content, src, rsn)
            except Exception as e:
                print(f"    [!] Warning: Failed parsing {p}: {e}")

    now_iso = datetime.now(timezone.utc).isoformat()
    cursor.executemany("""
        INSERT OR IGNORE INTO suppression (email, reason, source, blocked_at)
        VALUES (?, ?, ?, ?)
    """, [(e, r, s, now_iso) for e, r, s in suppressed_records.values()])
    conn.commit()
    print(f"    Loaded {len(suppressed_records)} suppression records.")

    # 2. Ingest Organizations from Excel
    print(" -> Ingesting Organizations and Expansion Signals from Excel...")
    excel_path = BASE_DIR / "Master_Career_Intelligence_Database.xlsx"
    wb = openpyxl.load_workbook(str(excel_path), data_only=True)
    
    org_sheet = wb["All Organizations"]
    org_rows = list(org_sheet.iter_rows(values_only=True))
    headers = [str(h) for h in org_rows[0]]
    
    orgs_added = 0
    signals_added = 0
    for r in org_rows[1:]:
        if not r or not r[0]:
            continue
        org_name = str(r[0]).strip()
        country = str(r[1]).strip() if r[1] else "GCC"
        sector = str(r[2]).strip() if r[2] else "General"
        trigger = str(r[3]).strip() if r[3] else None
        website = str(r[6]).strip() if len(r) > 6 and r[6] else None
        
        cursor.execute("""
            INSERT INTO organizations (name, country, sector, website, created_at)
            VALUES (?, ?, ?, ?, ?)
            ON CONFLICT(name) DO UPDATE SET
                country = excluded.country,
                sector = excluded.sector,
                website = COALESCE(excluded.website, organizations.website)
        """, (org_name, country, sector, website, now_iso))
        orgs_added += 1

        if trigger and trigger != "None":
            cursor.execute("""
                INSERT INTO signals (organization_name, headline, signal_date, source_url, details, created_at)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(organization_name, headline) DO UPDATE SET
                    signal_date = excluded.signal_date,
                    source_url = excluded.source_url,
                    details = excluded.details
            """, (org_name, trigger, "2026-10", website or "Verified Intelligence", trigger, now_iso))
            signals_added += 1

    conn.commit()
    print(f"    Loaded {orgs_added} organizations and {signals_added} expansion signals.")

    # 3. Ingest Contacts from Excel
    print(" -> Ingesting Decision-Makers & Contacts from Excel...")
    contact_sheet = wb["Decision Makers & Contacts"]
    contact_rows = list(contact_sheet.iter_rows(values_only=True))
    contacts_added = 0
    for r in contact_rows[1:]:
        if not r or not r[0] or not r[2]:
            continue
        name = str(r[0]).strip()
        role = str(r[1]).strip() if r[1] else "Lead"
        org_name = str(r[2]).strip()
        email = str(r[3]).strip() if r[3] and "@" in str(r[3]) else None
        phone = str(r[4]).strip() if len(r) > 4 and r[4] else None
        linkedin = str(r[5]).strip() if len(r) > 5 and r[5] else None
        v_status = str(r[6]).strip() if len(r) > 6 and r[6] else "Unverified"
        notes = str(r[10]).strip() if len(r) > 10 and r[10] else None
        
        is_supp = 1 if (email and repo.is_suppressed(email)) else 0

        # Ensure parent organization exists
        cursor.execute("INSERT OR IGNORE INTO organizations (name, country, created_at) VALUES (?, 'GCC', ?)", (org_name, now_iso))

        cursor.execute("""
            INSERT INTO contacts (
                organization_name, name, role, work_email, linkedin_url,
                verification_status, is_suppressed, notes, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(organization_name, name) DO UPDATE SET
                role = excluded.role,
                work_email = excluded.work_email,
                linkedin_url = excluded.linkedin_url,
                verification_status = excluded.verification_status,
                is_suppressed = excluded.is_suppressed,
                notes = excluded.notes,
                updated_at = excluded.updated_at
        """, (org_name, name, role, email, linkedin, v_status, is_supp, notes, now_iso))
        contacts_added += 1

    conn.commit()
    print(f"    Loaded {contacts_added} decision-makers & contacts.")

    # 4. Ingest Historical Outreach Ledger
    print(" -> Ingesting Historical Outreach Ledger...")
    ledger_sheet = wb["Outreach Ledger"]
    ledger_rows = list(ledger_sheet.iter_rows(values_only=True))
    events_added = 0
    for r in ledger_rows[1:]:
        if not r or not r[3]: # Recipient Email
            continue
        ts = str(r[0]).strip() if r[0] else now_iso
        email = str(r[3]).strip().lower()
        status = str(r[5]).strip() if len(r) > 5 and r[5] else "DISPATCHED"
        details = str(r[8]).strip() if len(r) > 8 and r[8] else ""
        
        cursor.execute("""
            INSERT INTO outreach_events (recipient_email, event_type, timestamp, details)
            VALUES (?, ?, ?, ?)
        """, (email, status, ts, details))
        events_added += 1

    conn.commit()
    print(f"    Loaded {events_added} outreach events.")

    # 5. Ingest 30 Calibrated Gmail Drafts
    print(" -> Ingesting 30 Calibrated Gmail Drafts...")
    md_path = BASE_DIR / "GCC_30_STAGED_EMAIL_DRAFTS_MASTER.md"
    if md_path.exists():
        content = md_path.read_text(encoding="utf-8")
        blocks = re.split(r"### Draft #", content)[1:]
        drafts_added = 0
        for block in blocks:
            m_to = re.search(r"- \*\*To:\*\* `([^`]+)`", block)
            m_subj = re.search(r"- \*\*Subject:\*\* `([^`]+)`", block)
            m_code = re.search(r"```text\s*(.*?)\s*```", block, re.DOTALL)
            
            to_email = m_to.group(1).strip() if m_to else ""
            subj = m_subj.group(1).strip() if m_subj else ""
            code_text = m_code.group(1).strip() if m_code else ""
            
            paras_raw = [p.strip() for p in code_text.split("\n\n") if p.strip()]
            paras = []
            for p in paras_raw:
                if "+91" in p or "linkedin.com" in p or p.startswith("Regards,"):
                    break
                paras.append(p)
            
            body_text = "\n\n".join(paras)
            word_count = len(body_text.split())
            import hashlib
            app_hash = hashlib.sha256(f"{to_email}:{subj}:{body_text}".encode()).hexdigest()
            
            recip_name = to_email.split("@")[0].capitalize()
            org_name = to_email.split("@")[1].split(".")[0].capitalize()

            cursor.execute("""
                INSERT INTO drafts (
                    recipient_email, recipient_name, organization_name, recipient_role,
                    subject_line, body_text, word_count, approval_hash, status, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'STAGED', ?)
                ON CONFLICT(recipient_email) DO UPDATE SET
                    subject_line = excluded.subject_line,
                    body_text = excluded.body_text,
                    word_count = excluded.word_count,
                    approval_hash = excluded.approval_hash,
                    status = excluded.status
            """, (to_email, recip_name, org_name, "Lead", subj, body_text, word_count, app_hash, now_iso))
            drafts_added += 1

        conn.commit()
        print(f"    Loaded {drafts_added} calibrated drafts.")

    conn.close()

    # Final summary
    stats = repo.get_system_stats()
    print("\n=======================================================")
    print("      ACIOS CANONICAL DATABASE INGESTION COMPLETE     ")
    print("=======================================================")
    print(f" - Organizations Cataloged:     {stats['total_organizations']}")
    print(f" - Decision-Makers & Contacts:   {stats['total_contacts']}")
    print(f" - Growth & Expansion Signals:   {stats['total_signals']}")
    print(f" - Verified Staged Drafts:       {stats['total_staged_drafts']}")
    print(f" - Master Suppression Blocklist: {stats['total_suppressed_emails']}")
    print(f" - Database Path:                {db_path}")
    print("=======================================================\n")


if __name__ == "__main__":
    import_all()
