"""
Unit tests for ACIOS Canonical Data Exporter.
"""

import sys
from pathlib import Path
import pytest
import openpyxl

PACKAGE_ROOT = Path(__file__).resolve().parent.parent
if str(PACKAGE_ROOT) not in sys.path:
    sys.path.insert(0, str(PACKAGE_ROOT))

from acios.core.database import DEFAULT_DB_PATH
from acios.reports.export import export_excel, export_drafts_markdown


def test_export_excel(tmp_path):
    output_excel = tmp_path / "Test_Master_Database.xlsx"
    export_excel(db_path=DEFAULT_DB_PATH, output_path=output_excel)

    assert output_excel.exists()
    wb = openpyxl.load_workbook(output_excel, data_only=True)
    assert "Executive Summary" in wb.sheetnames
    assert "All Organizations" in wb.sheetnames
    assert "Decision Makers & Contacts" in wb.sheetnames
    assert "Outreach Ledger" in wb.sheetnames

    # Check rows
    ws_orgs = wb["All Organizations"]
    assert ws_orgs.max_row >= 320


def test_export_drafts_markdown(tmp_path):
    output_md = tmp_path / "Test_Drafts.md"
    export_drafts_markdown(db_path=DEFAULT_DB_PATH, output_path=output_md)

    assert output_md.exists()
    content = output_md.read_text(encoding="utf-8")
    assert "### Draft #" in content
    assert "**Total Verified Staged Drafts:**" in content
